import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from document_selection import (
    DOCUMENT_CATEGORY_BY_TITLE,
    DOCUMENT_OVERVIEW_GROUPS,
    document_category_for_title,
    document_count_display_title,
    document_description_for_title,
    document_sort_key,
    normalize_document_title,
)


BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_DIR = BASE_DIR / "reference_templates"
COMPLIANCE_PLAN_TEMPLATE = TEMPLATE_DIR / "compliance_plan.docx"
DOCUMENT_LIST_TEMPLATE = TEMPLATE_DIR / "document_list.docx"


def _extract_document_name(text: str) -> str:
    match = re.search(r"《[^》]+》", text)
    if match:
        return match.group(0)
    return text.strip()


def _normalize_markdown(text: str) -> str:
    text = re.sub(r"^#{1,6}\s*", "", text.strip())
    return text.replace("```", "").strip()


def _split_table_row(text: str) -> list[str]:
    if "|" not in text:
        return []
    cells = [cell.strip() for cell in text.strip().strip("|").split("|")]
    if len(cells) < 2:
        return []
    non_empty = [cell.replace(" ", "") for cell in cells if cell]
    if non_empty and all(re.fullmatch(r":?-{3,}:?", cell) for cell in non_empty):
        return []
    return cells


def _clear_body(doc: Document) -> None:
    body = doc._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def _append_element(doc: Document, element) -> None:
    body = doc._body._element
    sect_pr = body.sectPr
    if sect_pr is None:
        body.append(element)
    else:
        body.insert(body.index(sect_pr), element)


def _replace_text_in_element(element, text: str) -> None:
    text_nodes = element.xpath(".//w:t")
    if not text_nodes:
        return
    text_nodes[0].text = text
    for node in text_nodes[1:]:
        node.text = ""


def _replace_cell_text(cell_element, text: str) -> None:
    text_nodes = cell_element.xpath(".//w:t")
    if text_nodes:
        text_nodes[0].text = text
        for node in text_nodes[1:]:
            node.text = ""
        return

    paragraphs = cell_element.xpath(".//w:p")
    paragraph = paragraphs[0] if paragraphs else OxmlElement("w:p")
    if not paragraphs:
        cell_element.append(paragraph)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    paragraph.append(run)


def _replace_row_cells(row_element, cells: list[str]) -> None:
    cell_elements = row_element.xpath("./w:tc")
    for index, cell_element in enumerate(cell_elements):
        _replace_cell_text(cell_element, cells[index] if index < len(cells) else "")


def _clone_paragraph(doc: Document, paragraph, text: str) -> None:
    element = copy.deepcopy(paragraph._p)
    _replace_text_in_element(element, text)
    _append_element(doc, element)


def _apply_compliance_body_style(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(20)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    for run in paragraph.runs:
        _set_font(run, "微软雅黑", 10, bold=False)
        _set_run_color(run, "3A4A5C")


def _clone_compliance_paragraph(doc: Document, paragraph, text: str, normalize_body: bool = False) -> None:
    _clone_paragraph(doc, paragraph, text)
    if normalize_body:
        _apply_compliance_body_style(doc.paragraphs[-1])


def _clone_section_table(doc: Document, table, text: str) -> None:
    element = copy.deepcopy(table._tbl)
    # 模板的分节条在第二列，第一列为空，用替换文本能保留原表格样式。
    text_nodes = element.xpath(".//w:t")
    if text_nodes:
        if len(text_nodes) >= 2:
            text_nodes[1].text = text
            for node in text_nodes[2:]:
                node.text = ""
        else:
            text_nodes[0].text = text
    _append_element(doc, element)


def _clone_section_table_by_index(doc: Document, tables: list, section_index: int) -> None:
    if 0 <= section_index < len(tables):
        _append_element(doc, copy.deepcopy(tables[section_index]._tbl))


def _set_font(run, name: str, size_pt: float, bold: bool = False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), name)


def _set_run_color(run, color: str) -> None:
    r_pr = run._element.get_or_add_rPr()
    color_el = r_pr.find(qn("w:color"))
    if color_el is None:
        color_el = OxmlElement("w:color")
        r_pr.append(color_el)
    color_el.set(qn("w:val"), color)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _set_line_spacing(paragraph, spacing_pt: float):
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:line"), str(int(spacing_pt * 20)))
    spacing.set(qn("w:lineRule"), "exact")
    p_pr.append(spacing)


def _add_runs(paragraph, text: str, font: str, size: float, bold: bool = False):
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        clean = part[2:-2] if is_bold else part
        run = paragraph.add_run(clean)
        _set_font(run, font, size, bold=(bold or is_bold))


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=col_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_index, row_cells in enumerate(rows):
        row = table.add_row()
        for cell_index, cell in enumerate(row.cells):
            value = row_cells[cell_index] if cell_index < len(row_cells) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            _set_line_spacing(paragraph, 16)
            _add_runs(paragraph, value, "仿宋", 10.5, bold=(row_index == 0))


def _is_category_row(cells: list[str]) -> bool:
    first = cells[0].strip() if cells else ""
    second = cells[1].strip() if len(cells) > 1 else ""
    if re.match(r"^[一二三四五六七八九十]+[、.．]", first) or re.match(r"^[一二三四五六七八九十]+[、.．]", second):
        return True
    if first and not first.isdigit() and first not in {"序号", "No"}:
        return True
    return False


def _document_title(cells: list[str]) -> str:
    value = cells[1] if len(cells) > 1 else cells[0] if cells else ""
    return normalize_document_title(_extract_document_name(value))


def _category_title(cells: list[str]) -> str:
    for cell in cells:
        text = cell.strip()
        if re.match(r"^[一二三四五六七八九十]+[、.．]", text):
            return text
    return cells[0].strip() if cells else ""


def _rebuild_document_overview_rows(rows: list[list[str]], document_counts: dict[str, int] | None = None) -> list[list[str]]:
    if not rows:
        return rows

    header = rows[0]
    selected_titles: set[str] = set()
    unknown_rows: list[list[str]] = []
    count_categories = getattr(document_counts, "categories", {}) if document_counts else {}
    for row in rows[1:]:
        if _is_category_row(row):
            continue
        title = _document_title(row)
        display_title = document_count_display_title(title)
        if document_category_for_title(display_title, count_categories):
            selected_titles.add(display_title)
        elif title:
            unknown_rows.append(row)

    if document_counts:
        selected_titles.update(title for title in document_counts)

    if not selected_titles:
        return rows

    rebuilt_rows = [header]
    sequence = 1
    used_titles: set[str] = set()
    for category, titles in DOCUMENT_OVERVIEW_GROUPS:
        category_titles = [
            title
            for title in selected_titles
            if document_category_for_title(title, count_categories) == category
        ]
        category_titles.sort(key=lambda item: document_sort_key(item, count_categories))
        if not category_titles:
            continue
        rebuilt_rows.append([category, category, ""])
        for title in category_titles:
            rebuilt_rows.append([str(sequence), f"《{title}》", ""])
            sequence += 1
            used_titles.add(title)

    remaining_titles = sorted(title for title in selected_titles if title not in used_titles)
    if remaining_titles:
        rebuilt_rows.append(["其他动态文书", "其他动态文书", ""])
        for title in remaining_titles:
            rebuilt_rows.append([str(sequence), f"《{title}》", ""])
            sequence += 1

    rebuilt_rows.extend(unknown_rows)
    return rebuilt_rows


def _add_document_list_overview_table(
    doc: Document,
    template_table,
    rows: list[list[str]],
    document_counts: dict[str, int] | None = None,
) -> None:
    if not rows:
        return

    document_counts = document_counts or {}
    rows = _rebuild_document_overview_rows(rows, document_counts)
    rendered_rows: list[list[str]] = []
    total_count = 0
    for row_index, row_cells in enumerate(rows):
        normalized_row = list(row_cells)
        if row_index > 0 and not _is_category_row(normalized_row):
            document_name = _extract_document_name(normalized_row[1] if len(normalized_row) > 1 else normalized_row[0])
            document_title = normalize_document_title(document_name)
            if document_title in document_counts:
                template_count = document_counts[document_title]
            elif document_counts:
                template_count = 1
            else:
                template_count = 1
            total_count += template_count
            while len(normalized_row) < 3:
                normalized_row.append("")
            normalized_row[2] = str(template_count)
        rendered_rows.append(normalized_row)

    if total_count:
        rendered_rows.append(["总计", "总计", str(total_count)])

    table_element = copy.deepcopy(template_table._tbl)
    for row in list(table_element.xpath("./w:tr")):
        table_element.remove(row)

    header_template = template_table.rows[0]._tr
    category_template = template_table.rows[1]._tr if len(template_table.rows) > 1 else template_table.rows[0]._tr
    item_template = template_table.rows[2]._tr if len(template_table.rows) > 2 else template_table.rows[0]._tr

    for row_index, row_cells in enumerate(rendered_rows):
        if row_index == 0:
            row_element = copy.deepcopy(header_template)
            normalized = ["序号", "文书名称", "份数"]
        elif row_cells and row_cells[0] == "总计":
            row_element = copy.deepcopy(item_template)
            normalized = row_cells
        elif _is_category_row(row_cells):
            row_element = copy.deepcopy(category_template)
            category_text = next((cell for cell in row_cells if re.match(r"^[一二三四五六七八九十]+[、.．]", cell.strip())), row_cells[0])
            physical_cell_count = len(row_element.xpath("./w:tc"))
            normalized = [category_text, ""] if physical_cell_count == 2 else [category_text, category_text, ""]
        else:
            row_element = copy.deepcopy(item_template)
            normalized = list(row_cells)
        _replace_row_cells(row_element, normalized)
        table_element.append(row_element)

    _append_element(doc, table_element)
    table = doc.tables[-1]
    for row_index, row in enumerate(table.rows):
        row_cells = [cell.text.strip() for cell in row.cells]
        is_header = row_index == 0
        is_total = row_cells and row_cells[0] == "总计"
        is_category = row_index > 0 and (_is_category_row(row_cells) or is_total)
        for cell_index, cell in enumerate(row.cells):
            if is_header:
                _set_cell_shading(cell, "1F497D")
            elif is_category:
                _set_cell_shading(cell, "D9E2F3")

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if (is_header or cell_index == 0) else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    if is_header:
                        _set_font(run, "黑体", 11, bold=True)
                        _set_run_color(run, "FFFFFF")
                    elif is_category:
                        _set_font(run, "黑体", 10.5, bold=True)
                        _set_run_color(run, "1F497D")
                    else:
                        _set_font(run, "仿宋", 10.5, bold=(cell_index == 1))
                        _set_run_color(run, "000000")


def _is_compliance_section(text: str) -> bool:
    return text.startswith(("第一部分", "第二部分", "第三部分"))


def _is_doclist_table_start(text: str) -> bool:
    cells = _split_table_row(text)
    return len(cells) >= 3 and cells[0] == "序号" and "文书名称" in cells[1]


def _is_compliance_body_paragraph(paragraph) -> bool:
    text = paragraph.text.strip()
    if len(text) <= 40:
        return False
    if text.startswith("· ·"):
        return False
    if text.startswith(("第一部分", "第二部分", "第三部分")):
        return False
    if text.startswith(("（", "整改要求：", "适用对象：", "核心要点：", "配套动作：")):
        return False
    if re.match(r"^\d+[\.．]", text):
        return False
    return True


def _pick_compliance_paragraph(template: Document, text: str):
    paragraphs = [p for p in template.paragraphs if p.text.strip()]
    cover_company = paragraphs[0]
    cover_en = paragraphs[1]
    cover_dots = paragraphs[2]
    cover_title = paragraphs[3]
    cover_line = paragraphs[4]
    h2 = next(p for p in paragraphs if p.text.strip().startswith("（一）"))
    body = next(p for p in paragraphs if _is_compliance_body_paragraph(p))

    if text == "ENTERPRISE LABOR COMPLIANCE PLAN":
        return cover_en
    if text == "劳动用工合规计划书":
        return cover_title
    if text.startswith("· ·"):
        return cover_dots
    if text.startswith("━"):
        return cover_line
    if text.startswith("（") or re.match(r"^\d+[\.．]", text):
        return h2
    if "公司" in text and len(text) <= 40:
        return cover_company
    return body


def _compliance_paragraph_styles(template: Document) -> dict[str, object]:
    paragraphs = [p for p in template.paragraphs if p.text.strip()]
    return {
        "cover_company": paragraphs[0],
        "cover_en": paragraphs[1],
        "cover_dots": paragraphs[2],
        "cover_title": paragraphs[3],
        "cover_line": paragraphs[4],
        "h2": next(p for p in paragraphs if p.text.strip().startswith("（一）")),
        "item_title": next(p for p in paragraphs if re.match(r"^\d+[\.．]", p.text.strip())),
        "requirement": next(p for p in paragraphs if p.text.strip().startswith("整改要求：")),
        "target": next(p for p in paragraphs if p.text.strip().startswith("适用对象：")),
        "key": next(p for p in paragraphs if p.text.strip().startswith("核心要点：")),
        "action": next(p for p in paragraphs if p.text.strip().startswith("配套动作：")),
        "body": next(p for p in paragraphs if _is_compliance_body_paragraph(p)),
    }


def _guess_compliance_part(line: str, current_part: int) -> int:
    if _is_compliance_section(line):
        if line.startswith("第一部分"):
            return 1
        if line.startswith("第二部分"):
            return 2
        if line.startswith("第三部分"):
            return 3

    if "整改" in line or line.startswith(("整改要求：", "配套动作：", "适用对象：", "核心要点：")):
        return max(current_part, 3)
    if line.startswith("（") and ("整改" in line or "建设" in line):
        return max(current_part, 3)
    if re.match(r"^\d+[\.．]", line) and current_part >= 2:
        return 3
    if "风险" in line or "管理风险" in line:
        return max(current_part, 2)
    if current_part == 0:
        return 1
    return current_part


def _is_compliance_cover_line(line: str) -> bool:
    return (
        line == "ENTERPRISE LABOR COMPLIANCE PLAN"
        or line == "劳动用工合规计划书"
        or line.startswith("· ·")
        or line.startswith("━")
    )


def _is_compliance_company_title(line: str, current_part: int) -> bool:
    return current_part == 0 and "公司" in line and len(line) <= 40


def _is_repeated_section_caption(line: str) -> bool:
    compact = re.sub(r"\s+", "", line)
    return bool(
        re.match(r"^[一二三四五六七八九十]+[、.．]", compact)
        and any(
            title in compact
            for title in ("企业用工基本情况", "现有用工管理核心情况及风险", "劳动用工合规整改清单")
        )
    )


def _style_for_compliance_line(styles: dict[str, object], line: str, part: int):
    if line == "ENTERPRISE LABOR COMPLIANCE PLAN":
        return styles["cover_en"]
    if line == "劳动用工合规计划书":
        return styles["cover_title"]
    if line.startswith("· ·"):
        return styles["cover_dots"]
    if line.startswith("━"):
        return styles["cover_line"]
    if line.startswith("整改要求："):
        return styles["requirement"]
    if line.startswith("适用对象："):
        return styles["target"]
    if line.startswith("核心要点："):
        return styles["key"]
    if line.startswith("配套动作："):
        return styles["action"]
    if _is_compliance_company_title(line, part):
        return styles["cover_company"]
    if line.startswith("（"):
        return styles["h2"]
    if re.match(r"^\d+[\.．]", line):
        return styles["item_title"] if part >= 3 else styles["h2"]
    return styles["body"]


def _pick_doclist_paragraph(template: Document, text: str):
    paragraphs = [p for p in template.paragraphs if p.text.strip()]
    title = paragraphs[0]
    overview = paragraphs[1]
    desc_title = next(p for p in paragraphs if p.text.strip() == "各文书作用说明")
    h1 = next(p for p in paragraphs if p.text.strip().startswith("一、"))
    h2 = next(p for p in paragraphs if re.match(r"^\d+[\.．]", p.text.strip()))
    body = next(p for p in paragraphs if len(p.text.strip()) > 40)

    if text == "劳动用工合规专项服务文书清单":
        return title
    if text.startswith("▌"):
        return overview
    if text == "各文书作用说明":
        return desc_title
    if re.match(r"^[一二三四五六七八九十]+、", text):
        return h1
    if re.match(r"^\d+[\.．]", text):
        return h2
    return body


def _doclist_paragraph_styles(template: Document) -> dict[str, object]:
    paragraphs = [p for p in template.paragraphs if p.text.strip()]
    return {
        "title": paragraphs[0],
        "overview": paragraphs[1],
        "desc_title": next(p for p in paragraphs if p.text.strip() == "各文书作用说明"),
        "h1": next(p for p in paragraphs if p.text.strip().startswith("一、")),
        "item_title": next(p for p in paragraphs if re.match(r"^\d+[\.．]", p.text.strip())),
        "body": next(
            p
            for p in paragraphs
            if len(p.text.strip()) > 40 and not re.match(r"^\d+[\.．]", p.text.strip())
        ),
    }


def _split_doclist_item_explanation(line: str) -> tuple[str, str] | None:
    match = re.match(r"^(\d+[\.．]\s*.+?)(?:\s*[：:]\s*)(.+)$", line)
    if not match:
        return None
    return match.group(1).strip(), match.group(2).strip()


def _format_standalone_doclist_item_title(line: str, item_number: int) -> str | None:
    text = re.sub(r"^[\-•·]\s*", "", line.strip())
    match = re.fullmatch(r"[（(]?\s*(《[^》]+》)\s*[）)]?", text)
    if not match:
        return None
    return f"{item_number}. {match.group(1)}"


def _is_doclist_explanation_heading(line: str) -> bool:
    return "各文书作用说明" in re.sub(r"\s+", "", line)


def _render_compliance_plan(text: str, output_path: str) -> None:
    template = Document(COMPLIANCE_PLAN_TEMPLATE)
    doc = Document(COMPLIANCE_PLAN_TEMPLATE)
    _clear_body(doc)

    section_tables = list(template.tables)
    styles = _compliance_paragraph_styles(template)
    inserted_sections: set[int] = set()
    current_part = 0
    lines = [_normalize_markdown(raw) for raw in text.splitlines()]
    skip_indices: set[int] = set()

    en_index = next((i for i, line in enumerate(lines) if line == "ENTERPRISE LABOR COMPLIANCE PLAN"), None)
    title_index = next((i for i, line in enumerate(lines) if line == "劳动用工合规计划书"), None)
    company_index = None
    if en_index is not None:
        for i in range(en_index - 1, -1, -1):
            line = lines[i]
            if not line:
                continue
            if _is_compliance_cover_line(line) or _is_compliance_section(line) or _is_repeated_section_caption(line):
                continue
            if len(line) <= 40:
                company_index = i
                break

    # The cover must follow the reference template order even when the LLM emits it
    # after the first section heading.
    if company_index is not None:
        _clone_compliance_paragraph(doc, styles["cover_company"], lines[company_index])
        skip_indices.add(company_index)
    if en_index is not None:
        _clone_compliance_paragraph(doc, styles["cover_en"], lines[en_index])
        skip_indices.add(en_index)
    if styles["cover_dots"].text.strip():
        _clone_compliance_paragraph(doc, styles["cover_dots"], styles["cover_dots"].text.strip())
    if title_index is not None:
        _clone_compliance_paragraph(doc, styles["cover_title"], lines[title_index])
        skip_indices.add(title_index)

    for index, line in enumerate(lines):
        if index in skip_indices:
            continue
        if not line:
            doc.add_paragraph()
            continue

        if _is_repeated_section_caption(line):
            continue

        if _is_compliance_cover_line(line) or _is_compliance_company_title(line, current_part):
            _clone_compliance_paragraph(doc, _style_for_compliance_line(styles, line, current_part), line)
            continue

        new_part = _guess_compliance_part(line, current_part)
        if new_part and new_part != current_part:
            current_part = new_part
            if current_part not in inserted_sections:
                _clone_section_table_by_index(doc, section_tables, current_part - 1)
                inserted_sections.add(current_part)

        if _is_compliance_section(line):
            continue

        is_body = current_part in (1, 2) and not (
            line.startswith(("（", "整改要求：", "适用对象：", "核心要点：", "配套动作："))
            or re.match(r"^\d+[\.．]", line)
        )
        _clone_compliance_paragraph(
            doc,
            _style_for_compliance_line(styles, line, current_part),
            line,
            normalize_body=is_body,
        )

    doc.save(output_path)


def _extract_document_explanations(lines: list[str]) -> dict[str, str]:
    explanations: dict[str, str] = {}
    in_explanation_section = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if _is_doclist_explanation_heading(line):
            in_explanation_section = True
            i += 1
            continue
        if not in_explanation_section:
            i += 1
            continue
        if not line:
            i += 1
            continue
        if re.match(r"^[一二三四五六七八九十]+、", line):
            i += 1
            continue
        split_item = _split_doclist_item_explanation(line)
        if split_item:
            item_title, explanation = split_item
            title_match = re.search(r"《([^》]+)》", item_title)
            if title_match:
                explanations[title_match.group(1)] = explanation
            i += 1
            continue
        if re.match(r"^\d+[\.．]", line):
            title_match = re.search(r"《([^》]+)》", line)
            if title_match:
                explanations[title_match.group(1)] = ""
            i += 1
            continue
        standalone_match = re.search(r"《([^》]+)》", line)
        if standalone_match:
            explanations[standalone_match.group(1)] = ""
        i += 1
    return explanations


def _render_document_list(text: str, output_path: str, document_counts: dict[str, int] | None = None) -> None:
    template = Document(DOCUMENT_LIST_TEMPLATE)
    doc = Document(DOCUMENT_LIST_TEMPLATE)
    _clear_body(doc)

    styles = _doclist_paragraph_styles(template)
    lines = [_normalize_markdown(line) for line in text.splitlines()]
    explanations = _extract_document_explanations(lines)

    i = 0
    rendered_overview_table = False
    in_explanation_section = False
    doclist_item_number = 1
    while i < len(lines):
        line = lines[i]
        if not line:
            doc.add_paragraph()
            i += 1
            continue

        row = _split_table_row(line)
        if row:
            rows = [row]
            i += 1
            while i < len(lines):
                next_row = _split_table_row(lines[i])
                if not next_row:
                    break
                rows.append(next_row)
                i += 1
            if not rendered_overview_table and template.tables:
                _add_document_list_overview_table(doc, template.tables[0], rows, document_counts)
                rendered_overview_table = True
            else:
                _add_table(doc, rows)
            continue

        if _is_doclist_explanation_heading(line):
            in_explanation_section = True
            if rendered_overview_table:
                doc.add_page_break()
            _clone_paragraph(doc, styles["desc_title"], line)
            i += 1
            continue

        if in_explanation_section:
            _render_explanation_section(doc, document_counts, explanations, styles)
            in_explanation_section = False
            break

        _clone_paragraph(doc, _pick_doclist_paragraph(template, line), line)
        i += 1

    doc.save(output_path)


def _render_explanation_section(doc: Document, document_counts: dict[str, int] | None, explanations: dict[str, str], styles: dict) -> None:
    from document_selection import document_sort_key, document_category_for_title, DOCUMENT_OVERVIEW_GROUPS

    doclist_item_number = 1
    count_categories = getattr(document_counts, "categories", {}) if document_counts else {}
    rendered_titles: set[str] = set()

    for category, _ in DOCUMENT_OVERVIEW_GROUPS:
        category_titles = []
        if document_counts:
            for title in document_counts:
                if document_category_for_title(title, count_categories) == category:
                    category_titles.append(title)
        category_titles.sort(key=lambda item: document_sort_key(item, count_categories))

        if not category_titles:
            continue

        _clone_paragraph(doc, styles["h1"], category)

        for title in category_titles:
            item_title = f"{doclist_item_number}. 《{title}》"
            _clone_paragraph(doc, styles["item_title"], item_title)

            explanation_text = document_description_for_title(title) or explanations.get(title, "")
            if explanation_text:
                _clone_paragraph(doc, styles["body"], explanation_text)
            else:
                _clone_paragraph(doc, styles["body"], "建议结合企业实际确认是否启用")

            doclist_item_number += 1
            rendered_titles.add(title)

    rest_titles = [title for title in (document_counts or {}) if title not in rendered_titles]
    if rest_titles:
        _clone_paragraph(doc, styles["h1"], "其他动态文书")
        for title in sorted(rest_titles):
            item_title = f"{doclist_item_number}. 《{title}》"
            _clone_paragraph(doc, styles["item_title"], item_title)
            explanation_text = document_description_for_title(title) or explanations.get(title, "")
            _clone_paragraph(doc, styles["body"], explanation_text or "建议结合企业实际确认是否启用")
            doclist_item_number += 1


def text_to_docx(
    text: str,
    output_path: str,
    doc_index: int | None = None,
    document_counts: dict[str, int] | None = None,
) -> None:
    if doc_index == 1:
        _render_compliance_plan(text, output_path)
    elif doc_index == 2:
        _render_document_list(text, output_path, document_counts)
    else:
        # 兼容旧调用：没有指定 doc_index 时仍按合规计划书模板输出。
        _render_compliance_plan(text, output_path)
