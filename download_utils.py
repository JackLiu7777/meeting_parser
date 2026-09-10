from io import BytesIO

from docx import Document


def _decode_text_content(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("gbk", errors="replace")


def _extract_docx_content(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_lines = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_lines.append(" | ".join(cells))
    return "\n".join(paragraphs + table_lines)


def _guess_file_ext(filename_or_url: str, content_type: str = "") -> str:
    lower = filename_or_url.lower()
    if lower.endswith(".docx"):
        return ".docx"
    if lower.endswith(".txt"):
        return ".txt"
    if "wordprocessingml.document" in content_type:
        return ".docx"
    if content_type.startswith("text/"):
        return ".txt"
    return ""


def extract_meeting_content(content: bytes, filename_or_url: str = "", content_type: str = "") -> str:
    """从上传的字节内容中提取会议纪要纯文本，支持 .txt 和 .docx。"""
    ext = _guess_file_ext(filename_or_url, content_type)
    if ext == ".docx":
        return _extract_docx_content(content)
    if ext in ("", ".txt"):
        return _decode_text_content(content)
    raise ValueError("仅支持 .txt 和 .docx 格式的会议纪要")
